"""
Tests for the non-LLM parts of the pipeline (extraction + authenticity).
Gemini structuring is mocked/skipped here — test that separately with a
live key, since it costs an API call and isn't deterministic.
"""
import io
from unittest.mock import Mock
import pytest
from pydantic import ValidationError
from docx import Document

from app.services.resume_extractor import (
    extract_raw_text,
    extract_text_from_docx,
    extract_github_username,
    UnsupportedFileTypeError,
    EmptyResumeError,
    InvalidResumeFileError,
)
from app.services.authenticity_checker import check_authenticity
from app.models.resume_schema import AuthenticityResult
from app.services import resume_structurer
from app.services.gemini_client import GeminiAPIError


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_table_docx_bytes() -> bytes:
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table-based resume"
    table.cell(0, 1).text = "Backend engineer with Python and database experience"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf_bytes(text: str | None) -> bytes:
    stream = (f"BT /F1 12 Tf 20 260 Td ({text}) Tj ET" if text else "q Q").encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 500 300] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF".encode("ascii"))
    return bytes(pdf)


def test_extract_text_from_docx():
    content = _make_docx_bytes(["Jane Doe", "Python developer with 3 years experience"])
    text = extract_text_from_docx(content)
    assert "Jane Doe" in text
    assert "Python developer" in text


def test_extract_github_username():
    assert extract_github_username("Portfolio: https://github.com/torvalds and email") == "torvalds"
    assert extract_github_username("Check github.com/octocat/ for code") == "octocat"
    assert extract_github_username("See https://github.com/about or pricing") is None
    assert extract_github_username("No link here") is None


def test_extracts_resume_text_from_docx_tables():
    text = extract_raw_text(_make_table_docx_bytes(), "resume.docx")
    assert "Table-based resume" in text
    assert "Backend engineer" in text


def test_extracts_text_from_pdf():
    content = _make_pdf_bytes(
        "Senior software engineer with Python experience building reliable web services for teams"
    )
    text = extract_raw_text(content, "resume.pdf")
    assert "Senior software engineer" in text
    assert "reliable web services" in text


def test_valid_pdf_without_text_layer_is_reported_as_empty():
    with pytest.raises(EmptyResumeError, match="scanned image"):
        extract_raw_text(_make_pdf_bytes(None), "scan.pdf")


def test_extract_raw_text_rejects_unsupported_extension():
    with pytest.raises(UnsupportedFileTypeError):
        extract_raw_text(b"hello", "resume.txt")


@pytest.mark.parametrize(
    ("content", "filename"),
    [
        (b"not a PDF file", "resume.pdf"),
        (b"not a DOCX file", "resume.docx"),
        (b"%PDF-1.7\nnot a valid PDF structure", "resume.pdf"),
    ],
)
def test_extract_raw_text_rejects_mismatched_or_corrupt_file_signatures(content, filename):
    with pytest.raises(InvalidResumeFileError):
        extract_raw_text(content, filename)


def test_extract_raw_text_rejects_empty_content():
    tiny_docx = _make_docx_bytes(["hi"])
    with pytest.raises(EmptyResumeError):
        extract_raw_text(tiny_docx, "resume.docx")


def test_scanned_pdf_without_text_layer_is_reported_as_empty(monkeypatch):
    from app.services import resume_extractor

    monkeypatch.setattr(resume_extractor, "extract_text_from_pdf", lambda _content: "")
    with pytest.raises(EmptyResumeError, match="scanned image"):
        extract_raw_text(b"%PDF-1.7\nsynthetic scanned document", "scan.pdf")


def test_authenticity_clean_resume_scores_high():
    text = (
        "Built a real-time chat application using WebSockets and Redis. "
        "Led a team of four engineers to migrate a monolith to microservices. "
        "Implemented CI pipelines with GitHub Actions reducing deploy time by 40 percent."
    )
    result = check_authenticity(text)
    assert result.authenticityScore >= 80
    assert result.flags == []


def test_authenticity_flags_known_boilerplate():
    text = (
        "Results driven professional with proven track record of success. "
        "Passionate about leveraging technology to solve problems. "
        "Highly motivated self starter with extensive experience in delivering value."
    )
    result = check_authenticity(text)
    assert "ai_generated_pattern" in result.flags
    assert result.authenticityScore < 100


def test_authenticity_flags_internal_repetition():
    text = " ".join(["synergized cross functional teams to deliver results"] * 4)
    result = check_authenticity(text)
    assert "duplicate_content" in result.flags


@pytest.mark.parametrize("score", [-1, 101])
def test_authenticity_schema_rejects_out_of_range_score(score):
    with pytest.raises(ValidationError):
        AuthenticityResult(authenticityScore=score)


def test_resume_schema_rejects_wrong_field_types():
    from app.models.resume_schema import ResumeParsed

    with pytest.raises(ValidationError):
        ResumeParsed(skills=[{"name": "Python"}])


def test_resume_parsing_mode_defaults_to_gemini():
    from app.models.resume_schema import ResumeParseResponse, ResumeParsed

    response = ResumeParseResponse(
        rawText="Synthetic test resume text",
        parsed=ResumeParsed(skills=["Python"]),
        authenticity=AuthenticityResult(authenticityScore=100),
    )
    assert response.parsingMode == "gemini"


def test_resume_parser_uses_disclosed_heuristic_fallback_when_gemini_is_unavailable(monkeypatch):
    def unavailable(*_args, **_kwargs):
        raise GeminiAPIError("temporarily unavailable", status_code=503)

    monkeypatch.setattr(resume_structurer, "call_gemini_json", unavailable)
    parsed, mode = resume_structurer.structure_resume_with_mode(
        "Python and React developer with PostgreSQL experience"
    )
    assert mode == "heuristic_fallback"
    assert parsed.skills == ["python", "react", "postgresql"]
    assert parsed.education == []


def test_resume_parser_uses_disclosed_fallback_when_gemini_rejects_request(monkeypatch):
    def rejected_request(*_args, **_kwargs):
        raise GeminiAPIError("provider rejected request", status_code=400)

    monkeypatch.setattr(resume_structurer, "call_gemini_json", rejected_request)
    parsed, mode = resume_structurer.structure_resume_with_mode(
        "Python and React developer with PostgreSQL experience"
    )

    assert mode == "heuristic_fallback"
    assert parsed.skills == ["python", "react", "postgresql"]


def test_resume_parser_does_not_hide_invalid_model_errors(monkeypatch):
    def invalid_model(*_args, **_kwargs):
        raise GeminiAPIError("model not found", status_code=404)

    monkeypatch.setattr(resume_structurer, "call_gemini_json", invalid_model)
    with pytest.raises(GeminiAPIError):
        resume_structurer.structure_resume_with_mode("Python developer")


def test_resume_parser_retries_schema_validation_once_then_succeeds(monkeypatch):
    from app.models.resume_schema import ResumeParsed

    response = Mock(side_effect=[{"skills": [123]}, {"skills": ["python"]}])
    monkeypatch.setattr(resume_structurer, "call_gemini_json", response)
    parsed, mode = resume_structurer.structure_resume_with_mode("Synthetic Python developer resume")
    assert parsed == ResumeParsed(skills=["python"])
    assert mode == "gemini"
    assert response.call_count == 2


def test_resume_parser_uses_disclosed_fallback_after_repeated_invalid_gemini_output(monkeypatch):
    response = Mock(return_value={"skills": [123]})
    monkeypatch.setattr(resume_structurer, "call_gemini_json", response)
    parsed, mode = resume_structurer.structure_resume_with_mode("Python developer")
    assert mode == "heuristic_fallback"
    assert parsed.skills == ["python"]
    assert response.call_count == 2


@pytest.mark.parametrize(("provider_status", "error_code"), [
    (400, "GEMINI_REQUEST_REJECTED"),
    (401, "GEMINI_API_KEY_INVALID"),
    (403, "GEMINI_ACCESS_DENIED"),
    (404, "GEMINI_MODEL_NOT_FOUND"),
])
def test_resume_route_returns_specific_safe_gemini_error_code(monkeypatch, provider_status, error_code):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.routers import resume

    def unavailable_model(*_args, **_kwargs):
        raise GeminiAPIError("private provider response", status_code=provider_status)

    monkeypatch.setattr(resume, "structure_resume_with_mode", unavailable_model)
    response = TestClient(app).post(
        "/ai/resume/parse",
        files={"file": ("resume.docx", _make_docx_bytes(["Python developer with extensive experience building APIs and databases for several product teams across multiple projects"]), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 503
    assert response.headers["x-error-code"] == error_code
    assert "private provider response" not in response.text


def test_resume_route_returns_unprocessable_entity_for_signature_mismatch():
    from fastapi.testclient import TestClient
    from app.main import app

    response = TestClient(app).post(
        "/ai/resume/parse",
        files={"file": ("resume.pdf", b"not a PDF", "application/pdf")},
    )
    assert response.status_code == 422


def test_resume_route_returns_payload_too_large(monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    from app.routers import resume

    monkeypatch.setattr(resume.settings, "MAX_UPLOAD_MB", 0)
    response = TestClient(app).post(
        "/ai/resume/parse",
        files={"file": ("resume.docx", b"oversized", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 413
