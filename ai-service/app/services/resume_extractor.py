"""
Step 1 of the resume pipeline: turn an uploaded file into raw text.
Supports PDF (pdfplumber) and DOCX (python-docx). Anything else is rejected
before this is even called (see routers/resume.py).
"""
import io
import re
import zipfile
import pdfplumber
from docx import Document

GITHUB_USERNAME_REGEX = re.compile(
    r"(?:https?://)?(?:www\.)?github\.com/([A-Za-z0-9](?:[A-Za-z0-9-]{0,37}[A-Za-z0-9])?)",
    re.IGNORECASE,
)
RESERVED_GITHUB_PATHS = {
    "about", "pricing", "features", "topics", "collections", "security",
    "enterprise", "customer-stories", "readme", "settings", "explore",
    "contact", "organizations", "login", "join", "marketplace", "sponsors",
}


def extract_github_username(raw_text: str) -> str | None:
    """Extract first valid public GitHub username mentioned in the raw resume text."""
    if not isinstance(raw_text, str):
        return None
    for match in GITHUB_USERNAME_REGEX.finditer(raw_text):
        username = match.group(1).strip()
        if username.lower() not in RESERVED_GITHUB_PATHS:
            return username
    return None


class UnsupportedFileTypeError(Exception):
    pass


class EmptyResumeError(Exception):
    pass


class InvalidResumeFileError(Exception):
    pass


MAX_EXTRACTED_CHARS = 50_000


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        if len(pdf.pages) > 10:
            raise InvalidResumeFileError("Resume exceeds maximum allowed page count (10 pages).")
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    table_cells = [
        cell.text.strip()
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
        if cell.text.strip()
    ]
    combined = "\n".join(paragraphs + table_cells).strip()
    if len(combined) > MAX_EXTRACTED_CHARS:
        return combined[:MAX_EXTRACTED_CHARS]
    return combined


def extract_raw_text(file_bytes: bytes, filename: str) -> str:
    """
    Dispatches to the right extractor based on file extension.
    Raises EmptyResumeError if extraction yields nothing usable — this
    usually means the PDF is a scanned image with no text layer.
    """
    lower = filename.lower()
    if lower.endswith(".pdf"):
        if not file_bytes.startswith(b"%PDF-"):
            raise InvalidResumeFileError("File content does not match the PDF extension")
        try:
            text = extract_text_from_pdf(file_bytes)
        except Exception as exc:
            raise InvalidResumeFileError("The PDF file is corrupt or invalid") from exc
    elif lower.endswith(".docx"):
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                required_parts = {"[Content_Types].xml", "word/document.xml"}
                if not required_parts.issubset(archive.namelist()):
                    raise InvalidResumeFileError("File content does not match the DOCX format")
                total_uncompressed = sum(info.file_size for info in archive.infolist())
                if total_uncompressed > 25 * 1024 * 1024:
                    raise InvalidResumeFileError("The DOCX archive exceeds maximum uncompressed size limits")
        except (zipfile.BadZipFile, OSError) as exc:
            raise InvalidResumeFileError("The DOCX file is corrupt or invalid") from exc
        try:
            text = extract_text_from_docx(file_bytes)
        except Exception as exc:
            raise InvalidResumeFileError("The DOCX file is corrupt or invalid") from exc
    else:
        raise UnsupportedFileTypeError(f"Unsupported file type: {filename}")

    if not text or len(text) < 50:
        raise EmptyResumeError(
            "Could not extract readable text. The file may be a scanned "
            "image without a text layer, or corrupted."
        )
    return text
